"""Create sample input files used by the tests (docx with image+table, pptx, xlsx, images, md, txt, pdf)."""
from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image, ImageDraw

OUT = Path(__file__).resolve().parent / "samples"


def make_all(out: Path = OUT) -> dict[str, Path]:
    out.mkdir(parents=True, exist_ok=True)
    files: dict[str, Path] = {}

    # --- images ---------------------------------------------------------
    im = Image.new("RGB", (1200, 800), (90, 216, 255))
    d = ImageDraw.Draw(im)
    d.rectangle([40, 40, 1160, 760], outline=(10, 30, 60), width=12)
    d.ellipse([400, 200, 800, 600], fill=(255, 209, 102))
    p = out / "发票01-高铁票.jpg"
    im.save(p, quality=90)
    files["jpg"] = p
    # portrait JPEG with EXIF orientation 6 (stored landscape, displayed portrait)
    im2 = Image.new("RGB", (1000, 700), (240, 240, 240))
    d2 = ImageDraw.Draw(im2)
    d2.rectangle([20, 20, 300, 680], fill=(110, 242, 166))
    d2.text((350, 300), "TOP", fill=(0, 0, 0))
    exif = im2.getexif()
    exif[0x0112] = 6
    p = out / "发票03-出租车-竖拍.jpg"
    im2.save(p, quality=90, exif=exif.tobytes())
    files["jpg_exif"] = p
    im3 = Image.new("RGBA", (800, 1100), (0, 0, 0, 0))
    d3 = ImageDraw.Draw(im3)
    d3.rectangle([60, 60, 740, 1040], fill=(167, 139, 250, 255))
    p = out / "截图-透明.png"
    im3.save(p)
    files["png"] = p

    # --- docx -------------------------------------------------------------
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    doc = Document()
    doc.add_heading("酒店住宿发票", 1)
    para = doc.add_paragraph("入住日期：2026-08-12 至 2026-08-15，共 3 晚。")
    r = para.add_run(" 金额：¥1,280.00")
    r.bold = True
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    doc.add_paragraph("项目一", style="List Bullet")
    doc.add_paragraph("项目二", style="List Bullet")
    doc.add_paragraph("步骤一", style="List Number")
    doc.add_paragraph("步骤二", style="List Number")
    t = doc.add_table(rows=3, cols=3)
    t.style = "Table Grid"
    for i, h in enumerate(["项目", "数量", "金额"]):
        t.cell(0, i).text = h
    for r_ in range(1, 3):
        t.cell(r_, 0).text = f"房费 {r_}"
        t.cell(r_, 1).text = "1"
        t.cell(r_, 2).text = "¥640.00"
    doc.add_paragraph()
    doc.add_picture(str(files["jpg"]), width=Inches(3))
    doc.add_page_break()
    doc.add_heading("第二页 附件", 2)
    doc.add_paragraph("这是第二页的内容。" * 20)
    p = out / "发票02-酒店住宿.docx"
    doc.save(p)
    files["docx"] = p

    # --- pptx -------------------------------------------------------------
    from pptx import Presentation
    from pptx.util import Inches as PInches, Pt as PPt
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Q3 项目总结"
    s.placeholders[1].text = "OrbPDF 内置引擎测试"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "要点"
    tf = s2.placeholders[1].text_frame
    tf.text = "第一条要点"
    pp = tf.add_paragraph()
    pp.text = "第二条要点（子项）"
    pp.level = 1
    s2.shapes.add_picture(str(files["jpg"]), PInches(5.5), PInches(3), width=PInches(3))
    box = s2.shapes.add_shape(1, PInches(0.5), PInches(5.5), PInches(3), PInches(1))  # rectangle
    box.text_frame.text = "形状里的文字"
    p = out / "项目总结.pptx"
    prs.save(p)
    files["pptx"] = p

    # --- xlsx -------------------------------------------------------------
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "报销明细"
    ws.append(["日期", "项目", "金额", "备注"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FFDDEEFF")
    for i in range(1, 30):
        ws.append([f"2026-08-{i:02d}", f"项目 {i}", i * 12.5, "自动生成" if i % 3 else ""])
    ws.merge_cells("A32:B32")
    ws["A32"] = "合计"
    ws["C32"] = "=SUM(C2:C30)"
    ws.column_dimensions["B"].width = 18
    ws2 = wb.create_sheet("宽表")
    ws2.append([f"列{i}" for i in range(1, 26)])
    for r_ in range(10):
        ws2.append([r_ * i for i in range(1, 26)])
    p = out / "报销明细.xlsx"
    wb.save(p)
    files["xlsx"] = p

    # --- text ------------------------------------------------------------
    p = out / "说明.md"
    p.write_text("# 说明\n\n这是 **Markdown** 文件。\n\n| 列 A | 列 B |\n|---|---|\n| 1 | 2 |\n\n- 项目一\n- 项目二\n", encoding="utf-8")
    files["md"] = p
    p = out / "日志.txt"
    p.write_text("第一行\n    缩进第二行\n\t制表符\nEnglish line\n" * 30, encoding="utf-8")
    files["txt"] = p

    # --- pdf (5 pages) ------------------------------------------------------
    import pymupdf
    pdf = pymupdf.open()
    for i in range(5):
        pg = pdf.new_page()
        pg.insert_text((72, 100), f"餐饮电子发票 第 {i + 1} 页", fontsize=20, fontname="china-s")
        pg.draw_rect(pymupdf.Rect(72, 140, 500, 400), color=(0.2, 0.5, 0.8), width=2)
    p = out / "发票04-餐饮.pdf"
    pdf.set_toc([[1, "第一部分", 1], [1, "第二部分", 3]])
    pdf.save(p)
    pdf.close()
    files["pdf"] = p
    return files


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else OUT
    for k, v in make_all(out).items():
        print(k, v)
