"""Built-in Word (.docx) engine: python-docx -> HTML -> PyMuPDF Story.

Keeps: headings, paragraphs, bold/italic/underline/strike, colour, size, alignment, indents,
bullets & numbering, tables (borders, spans, shading), inline images, page breaks, page size & margins.
Skips: headers/footers, footnotes, floating layout, text effects.
"""
from __future__ import annotations

import re
from pathlib import Path

import pymupdf

from ...fileinfo import KIND_WORD
from ..base import ConvertError, ConvertOptions, Engine, EngineInfo, ProgressFn, CancelFn, report, check_cancel
from .story_utils import esc, html_to_pdf, css_color

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
_V = "urn:schemas-microsoft-com:vml"


def _q(ns: str, tag: str) -> str:
    return "{%s}%s" % (ns, tag)


class DocxEngine(Engine):
    id = "builtin-docx"
    name = "内置 Word 引擎"
    kinds = frozenset({KIND_WORD})
    fidelity = "basic"

    def probe(self) -> EngineInfo:
        return EngineInfo(self.id, self.name, True, "docx 基础排版", set(self.kinds), self.fidelity)

    def supports(self, src: Path) -> bool:
        return src.suffix.lower() in (".docx", ".dotx")

    def convert(self, src: Path, dst: Path, opts: ConvertOptions, progress: ProgressFn = None, cancel: CancelFn = None) -> list[str]:
        report(progress, 0.1, "解析 Word 文档")
        try:
            conv = _DocxToHtml(src)
            html, css, archive, size, margins, warnings = conv.run(cancel)
        except ConvertError:
            raise
        except Exception as e:
            raise ConvertError(f"无法解析 docx：{e}") from e
        report(progress, 0.5, "排版")
        html_to_pdf(html, dst, css=css, archive=archive, page_size=size, margins=margins, progress=progress, cancel=cancel)
        return warnings


class _DocxToHtml:
    def __init__(self, path: Path) -> None:
        from docx import Document  # noqa: WPS433
        self.path = path
        self.doc = Document(str(path))
        self.parts: list[str] = []
        self.warnings: list[str] = []
        self.archive = pymupdf.Archive()
        self.img_count = 0
        self.counters: dict[tuple[str, int], int] = {}
        self._num_cache: dict[tuple[str, int], str] = {}
        self._has_header_footer = False

    # ---- entry -------------------------------------------------------------
    def run(self, cancel: CancelFn):
        doc = self.doc
        body = doc.element.body
        size, margins = self._page_geometry()
        base_size = self._default_font_size()
        for child in body.iterchildren():
            check_cancel(cancel)
            tag = child.tag
            if tag == _q(_W, "p"):
                self.parts.append(self._paragraph(child))
            elif tag == _q(_W, "tbl"):
                self.parts.append(self._table(child))
            elif tag == _q(_W, "sdt"):
                # content controls (e.g. Word's own TOC) -> render inner paragraphs
                for p in child.iter(_q(_W, "p")):
                    self.parts.append(self._paragraph(p))
        try:
            for sec in doc.sections:
                if sec.header and any(p.text.strip() for p in sec.header.paragraphs) or sec.footer and any(p.text.strip() for p in sec.footer.paragraphs):
                    self._has_header_footer = True
                    break
        except Exception:
            pass
        if self._has_header_footer:
            self.warnings.append("页眉页脚未包含（内置引擎）")
        css = f"body{{font-size:{base_size}pt}} .cell p{{margin:0}} td{{border:0.75pt solid #777}} .noborder td{{border:0}}"
        return "\n".join(self.parts), css, self.archive, size, margins, self.warnings

    # ---- geometry ----------------------------------------------------------
    def _page_geometry(self):
        try:
            s = self.doc.sections[0]
            w = float(s.page_width.pt) if s.page_width else 595.28
            h = float(s.page_height.pt) if s.page_height else 841.89
            m = (float(s.left_margin.pt) if s.left_margin else 56, float(s.top_margin.pt) if s.top_margin else 56,
                 float(s.right_margin.pt) if s.right_margin else 56, float(s.bottom_margin.pt) if s.bottom_margin else 56)
            m = tuple(max(18.0, min(v, w / 3 if i % 2 == 0 else h / 3)) for i, v in enumerate(m))
            return (w, h), m
        except Exception:
            return (595.28, 841.89), (56, 56, 56, 56)

    def _default_font_size(self) -> float:
        try:
            st = self.doc.styles["Normal"]
            if st.font and st.font.size:
                return float(st.font.size.pt)
        except Exception:
            pass
        return 10.5

    # ---- paragraphs ---------------------------------------------------------
    def _heading_level(self, p) -> int:
        try:
            name = (p.style.name or "").strip()
        except Exception:
            return 0
        m = re.match(r"(?:Heading|标题)\s*(\d)", name, re.I)
        if m:
            return min(6, int(m.group(1)))
        if name.lower() in ("title", "标题"):
            return 1
        return 0

    def _style_numpr(self, p_el):
        """numPr inherited from the paragraph style chain (e.g. 'List Bullet')."""
        try:
            pPr = p_el.find(_q(_W, "pPr"))
            style_id = None
            if pPr is not None:
                ps = pPr.find(_q(_W, "pStyle"))
                if ps is not None:
                    style_id = ps.get(_q(_W, "val"))
            if not style_id:
                return None
            styles_el = self.doc.styles.element
            seen = set()
            while style_id and style_id not in seen:
                seen.add(style_id)
                st = styles_el.find(f"{_q(_W, 'style')}[@{_q(_W, 'styleId')}='{style_id}']")
                if st is None:
                    return None
                spPr = st.find(_q(_W, "pPr"))
                if spPr is not None:
                    numPr = spPr.find(_q(_W, "numPr"))
                    if numPr is not None:
                        return numPr
                based = st.find(_q(_W, "basedOn"))
                style_id = based.get(_q(_W, "val")) if based is not None else None
        except Exception:
            return None
        return None

    def _numbering(self, p_el) -> str:
        pPr = p_el.find(_q(_W, "pPr"))
        numPr = pPr.find(_q(_W, "numPr")) if pPr is not None else None
        if numPr is None:
            numPr = self._style_numpr(p_el)
        if numPr is None:
            return ""
        numId_el = numPr.find(_q(_W, "numId"))
        ilvl_el = numPr.find(_q(_W, "ilvl"))
        if numId_el is None:
            return ""
        num_id = numId_el.get(_q(_W, "val"))
        ilvl = int(ilvl_el.get(_q(_W, "val"))) if ilvl_el is not None else 0
        if num_id in (None, "0"):
            return ""
        fmt = self._num_format(num_id, ilvl)
        if fmt == "bullet":
            return "•"
        key = (num_id, ilvl)
        self.counters[key] = self.counters.get(key, 0) + 1
        # reset deeper levels
        for k in list(self.counters):
            if k[0] == num_id and k[1] > ilvl:
                self.counters[k] = 0
        n = self.counters[key]
        if fmt in ("lowerLetter", "upperLetter"):
            s = chr(ord("a") + (n - 1) % 26)
            return (s.upper() if fmt == "upperLetter" else s) + "."
        if fmt in ("lowerRoman", "upperRoman"):
            s = _roman(n)
            return (s.upper() if fmt == "upperRoman" else s) + "."
        if fmt == "chineseCounting":
            return _chinese(n) + "、"
        return f"{n}."

    def _num_format(self, num_id: str, ilvl: int) -> str:
        key = (num_id, ilvl)
        if key in self._num_cache:
            return self._num_cache[key]
        fmt = "decimal"
        try:
            numbering = self.doc.part.numbering_part.element
            num = numbering.find(f"{_q(_W, 'num')}[@{_q(_W, 'numId')}='{num_id}']")
            if num is not None:
                abs_id = num.find(_q(_W, "abstractNumId")).get(_q(_W, "val"))
                absnum = numbering.find(f"{_q(_W, 'abstractNum')}[@{_q(_W, 'abstractNumId')}='{abs_id}']")
                if absnum is not None:
                    lvl = absnum.find(f"{_q(_W, 'lvl')}[@{_q(_W, 'ilvl')}='{ilvl}']")
                    if lvl is not None:
                        nf = lvl.find(_q(_W, "numFmt"))
                        if nf is not None:
                            fmt = nf.get(_q(_W, "val")) or "decimal"
        except Exception:
            pass
        self._num_cache[key] = fmt
        return fmt

    def _paragraph(self, p_el, in_cell: bool = False) -> str:
        from docx.text.paragraph import Paragraph  # noqa: WPS433
        p = Paragraph(p_el, self.doc)
        level = 0 if in_cell else self._heading_level(p)
        styles: list[str] = []
        try:
            al = p.alignment
            if al is not None:
                v = int(al)
                styles.append({1: "text-align:center", 2: "text-align:right", 3: "text-align:justify"}.get(v, ""))
        except Exception:
            pass
        try:
            pf = p.paragraph_format
            if pf.left_indent:
                styles.append(f"margin-left:{max(0.0, float(pf.left_indent.pt)):.1f}pt")
            if pf.first_line_indent and float(pf.first_line_indent.pt) > 0:
                styles.append(f"text-indent:{float(pf.first_line_indent.pt):.1f}pt")
            if pf.space_before:
                styles.append(f"margin-top:{float(pf.space_before.pt):.1f}pt")
            if pf.space_after is not None:
                styles.append(f"margin-bottom:{float(pf.space_after.pt):.1f}pt")
            if pf.line_spacing and isinstance(pf.line_spacing, float) and 0.8 <= pf.line_spacing <= 3:
                styles.append(f"line-height:{pf.line_spacing:.2f}")
        except Exception:
            pass
        prefix = "" if in_cell else self._numbering(p_el)
        segments: list[str] = []
        page_break_before = False
        try:
            pPr = p_el.find(_q(_W, "pPr"))
            if pPr is not None and pPr.find(_q(_W, "pageBreakBefore")) is not None:
                page_break_before = True
        except Exception:
            pass
        inner, breaks = self._runs_html(p)
        if prefix:
            inner = f"<span>{esc(prefix)}&nbsp;</span>" + inner
        tag = f"h{level}" if level else "p"
        style_attr = f' style="{";".join(s for s in styles if s)}"' if any(styles) else ""
        html = f"<{tag}{style_attr}>{inner or '&nbsp;'}</{tag}>"
        if breaks:
            html = html + '<div class="pb"></div>'
        if page_break_before:
            html = '<div class="pb"></div>' + html
        return html

    def _runs_html(self, p) -> tuple[str, bool]:
        out: list[str] = []
        page_break = False
        try:
            items = list(p.iter_inner_content())
        except Exception:
            items = list(p.runs)
        for item in items:
            runs = getattr(item, "runs", None)
            if runs is not None and not hasattr(item, "font"):
                # hyperlink
                inner = "".join(self._run_html(r) for r in runs)
                out.append(f'<span style="color:#1a56b0;text-decoration:underline">{inner}</span>')
            else:
                out.append(self._run_html(item))
            try:
                if item._r.findall(f".//{_q(_W, 'br')}[@{_q(_W, 'type')}='page']"):
                    page_break = True
            except Exception:
                pass
        return "".join(out), page_break

    def _run_html(self, run) -> str:
        r_el = run._r
        pieces: list[str] = []
        # images inside the run
        for blip in r_el.iter(_q(_A, "blip")):
            rid = blip.get(_q(_R, "embed"))
            if rid:
                pieces.append(self._image_html(rid, r_el))
        for imd in r_el.iter(_q(_V, "imagedata")):
            rid = imd.get(_q(_R, "id"))
            if rid:
                pieces.append(self._image_html(rid, r_el))
        text = run.text or ""
        if text:
            t = esc(text).replace("\t", "&nbsp;&nbsp;&nbsp;&nbsp;").replace("\n", "<br>")
            t = re.sub(r"  +", lambda m: "&nbsp;" * len(m.group(0)), t)
            styles: list[str] = []
            f = run.font
            try:
                if run.bold:
                    styles.append("font-weight:bold")
                if run.italic:
                    styles.append("font-style:italic")
                deco = []
                if run.underline:
                    deco.append("underline")
                if f.strike:
                    deco.append("line-through")
                if deco:
                    styles.append("text-decoration:" + " ".join(deco))
                if f.size:
                    styles.append(f"font-size:{float(f.size.pt):.1f}pt")
                c = None
                try:
                    c = css_color(f.color.rgb) if f.color and f.color.type is not None else None
                except Exception:
                    c = None
                if c and c.lower() != "#000000":
                    styles.append(f"color:{c}")
                try:
                    if f.highlight_color:
                        styles.append("background:#fff2a8")
                except Exception:
                    pass
                if f.superscript:
                    styles.append("vertical-align:super;font-size:70%")
                elif f.subscript:
                    styles.append("vertical-align:sub;font-size:70%")
            except Exception:
                pass
            if styles:
                pieces.append(f'<span style="{";".join(styles)}">{t}</span>')
            else:
                pieces.append(t)
        return "".join(pieces)

    def _image_html(self, rid: str, r_el) -> str:
        try:
            part = self.doc.part.related_parts[rid]
            blob = part.blob
            ct = (getattr(part, "content_type", "") or "").lower()
        except Exception:
            return ""
        ext = "png"
        if "jpeg" in ct or "jpg" in ct:
            ext = "jpg"
        elif "gif" in ct:
            ext = "gif"
        elif "bmp" in ct:
            ext = "bmp"
        elif "tiff" in ct:
            ext = "tif"
        elif "emf" in ct or "wmf" in ct or "svg" in ct:
            self.warnings.append("矢量图（EMF/WMF/SVG）已跳过")
            return ""
        self.img_count += 1
        name = f"img{self.img_count}.{ext}"
        try:
            self.archive.add(blob, name)
        except Exception:
            return ""
        w = h = None
        try:
            ext_el = next(r_el.iter(_q(_WP, "extent")), None)
            if ext_el is not None:
                w = float(ext_el.get("cx")) / 12700.0
                h = float(ext_el.get("cy")) / 12700.0
        except Exception:
            w = h = None
        attrs = f' width="{w:.1f}" height="{h:.1f}"' if w and h else ""
        return f'<img src="{name}"{attrs}>'

    # ---- tables -------------------------------------------------------------
    def _table(self, tbl_el) -> str:
        bordered = self._table_has_borders(tbl_el)
        rows_html: list[str] = []
        vmerge_skip: dict[int, int] = {}
        for tr in tbl_el.findall(_q(_W, "tr")):
            cells: list[str] = []
            col = 0
            for tc in tr.findall(_q(_W, "tc")):
                tcPr = tc.find(_q(_W, "tcPr"))
                span = 1
                vmerge = None
                width = None
                shade = None
                if tcPr is not None:
                    gs = tcPr.find(_q(_W, "gridSpan"))
                    if gs is not None:
                        try:
                            span = int(gs.get(_q(_W, "val")))
                        except Exception:
                            span = 1
                    vm = tcPr.find(_q(_W, "vMerge"))
                    if vm is not None:
                        vmerge = vm.get(_q(_W, "val")) or "continue"
                    tw = tcPr.find(_q(_W, "tcW"))
                    if tw is not None and tw.get(_q(_W, "type")) == "dxa":
                        try:
                            width = int(tw.get(_q(_W, "w"))) / 20.0
                        except Exception:
                            width = None
                    sh = tcPr.find(_q(_W, "shd"))
                    if sh is not None:
                        fill = sh.get(_q(_W, "fill"))
                        if fill and fill.lower() not in ("auto", "ffffff"):
                            shade = "#" + fill
                if vmerge == "continue":
                    col += span
                    continue
                inner = []
                for child in tc.iterchildren():
                    if child.tag == _q(_W, "p"):
                        inner.append(self._paragraph(child, in_cell=True))
                    elif child.tag == _q(_W, "tbl"):
                        inner.append(self._table(child))
                style = []
                if width:
                    style.append(f"width:{width:.1f}pt")
                if shade:
                    style.append(f"background:{shade}")
                attrs = f' colspan="{span}"' if span > 1 else ""
                if style:
                    attrs += f' style="{";".join(style)}"'
                cells.append(f'<td class="cell"{attrs}>{"".join(inner) or "&nbsp;"}</td>')
                col += span
            rows_html.append("<tr>" + "".join(cells) + "</tr>")
        cls = "" if bordered else ' class="noborder"'
        return f'<table{cls}>{"".join(rows_html)}</table><p style="margin:0 0 4pt 0"></p>'

    def _table_has_borders(self, tbl_el) -> bool:
        tblPr = tbl_el.find(_q(_W, "tblPr"))
        if tblPr is not None:
            b = tblPr.find(_q(_W, "tblBorders"))
            if b is not None:
                for side in b:
                    if (side.get(_q(_W, "val")) or "nil") not in ("nil", "none"):
                        return True
                return False
            st = tblPr.find(_q(_W, "tblStyle"))
            if st is not None:
                name = (st.get(_q(_W, "val")) or "").lower()
                if "grid" in name or "网格" in name or "table" in name:
                    return True
        # any cell-level borders?
        for tcb in tbl_el.iter(_q(_W, "tcBorders")):
            for side in tcb:
                if (side.get(_q(_W, "val")) or "nil") not in ("nil", "none"):
                    return True
        return False


def _roman(n: int) -> str:
    vals = [(1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"), (90, "xc"), (50, "l"), (40, "xl"), (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]
    out = ""
    for v, s in vals:
        while n >= v:
            out += s
            n -= v
    return out or "i"


def _chinese(n: int) -> str:
    digits = "零一二三四五六七八九"
    if n < 10:
        return digits[n]
    if n < 20:
        return "十" + (digits[n % 10] if n % 10 else "")
    if n < 100:
        return digits[n // 10] + "十" + (digits[n % 10] if n % 10 else "")
    return str(n)
