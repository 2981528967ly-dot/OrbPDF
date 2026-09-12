"""Built-in PDF -> Word (.docx): no Office needed.

Reads text blocks (font size / bold / italic / colour), detected tables and embedded images from each page
with PyMuPDF and rebuilds them as paragraphs, headings, tables and pictures with python-docx. Layout is
simplified (flowing text), content is complete. Word's own PDF reflow stays the high-fidelity option.
"""
from __future__ import annotations

import io
import statistics
from pathlib import Path

import pymupdf

from ..pdf.info import open_pdf
from .base import ConvertError, ProgressFn, CancelFn, report, check_cancel

_EMU_PER_PT = 12700


def _rgb(color_int: int) -> tuple[int, int, int]:
    return ((color_int >> 16) & 255, (color_int >> 8) & 255, color_int & 255)


def _is_bold(span: dict) -> bool:
    name = (span.get("font") or "").lower()
    return bool(span.get("flags", 0) & 16) or "bold" in name or "black" in name or "heavy" in name


def _is_italic(span: dict) -> bool:
    name = (span.get("font") or "").lower()
    return bool(span.get("flags", 0) & 2) or "italic" in name or "oblique" in name


def _join_lines(lines: list[str]) -> str:
    """Join wrapped lines: no space between CJK characters, a space between Latin words."""
    out = ""
    for ln in lines:
        ln = ln.strip()
        if not ln:
            continue
        if out and not (_cjk(out[-1]) or _cjk(ln[0])) and not out.endswith("-"):
            out += " "
        elif out.endswith("-") and ln and ln[0].islower():
            out = out[:-1]
        out += ln
    return out


def _cjk(ch: str) -> bool:
    return ord(ch) >= 0x2E80


def pdf_to_docx(src: Path, dst: Path, *, password: str | None = None, progress: ProgressFn = None, cancel: CancelFn = None) -> list[str]:
    try:
        from docx import Document  # noqa: WPS433
        from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: WPS433
        from docx.shared import Emu, Pt, RGBColor  # noqa: WPS433
    except Exception as e:  # pragma: no cover
        raise ConvertError(f"缺少 python-docx：{e}") from e
    doc = open_pdf(src, password)
    warnings: list[str] = []
    try:
        d = Document()
        first = doc[0].rect if doc.page_count else pymupdf.Rect(0, 0, 595, 842)
        sec = d.sections[0]
        sec.page_width = Emu(int(first.width * _EMU_PER_PT))
        sec.page_height = Emu(int(first.height * _EMU_PER_PT))
        margin = Emu(int(min(56.0, first.width * 0.08) * _EMU_PER_PT))
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = margin
        usable_w = first.width - 2 * min(56.0, first.width * 0.08)
        st = d.styles["Normal"]
        st.font.name = "Microsoft YaHei"
        try:
            from docx.oxml.ns import qn  # noqa: WPS433
            st.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        except Exception:
            pass
        n = doc.page_count
        n_tables = n_images = 0
        for pno in range(n):
            check_cancel(cancel)
            report(progress, pno / max(1, n), f"重建第 {pno + 1} 页")
            page = doc[pno]
            elements: list[tuple[float, float, str, object]] = []
            table_rects: list[pymupdf.Rect] = []
            try:
                tabs = page.find_tables()
                for tb in tabs.tables:
                    data = tb.extract()
                    if data and any(any(c for c in row) for row in data):
                        r = pymupdf.Rect(tb.bbox)
                        table_rects.append(r)
                        elements.append((r.y0, r.x0, "table", data))
            except Exception:
                pass
            try:
                info = page.get_text("dict")
            except Exception:
                info = {"blocks": []}
            sizes: list[float] = []
            for b in info.get("blocks", []):
                if b.get("type") == 0:
                    for ln in b.get("lines", []):
                        for sp in ln.get("spans", []):
                            if sp.get("text", "").strip():
                                sizes.append(float(sp.get("size", 11)))
            body_size = statistics.median(sizes) if sizes else 11.0
            for b in info.get("blocks", []):
                r = pymupdf.Rect(b["bbox"])
                if b.get("type") == 1:
                    elements.append((r.y0, r.x0, "image", b))
                    continue
                inside = False
                for tr in table_rects:
                    inter = r & tr
                    if not inter.is_empty and inter.get_area() > 0.5 * max(1e-6, r.get_area()):
                        inside = True
                        break
                if inside:
                    continue
                elements.append((r.y0, r.x0, "text", b))
            elements.sort(key=lambda e: (round(e[0] / 6.0), e[1]))
            for _y, _x, kind, payload in elements:
                if kind == "text":
                    self_lines = payload.get("lines", [])
                    spans_all = [sp for ln in self_lines for sp in ln.get("spans", []) if sp.get("text", "")]
                    if not spans_all:
                        continue
                    max_size = max(float(sp.get("size", body_size)) for sp in spans_all)
                    text_lines = ["".join(sp.get("text", "") for sp in ln.get("spans", [])) for ln in self_lines]
                    joined = _join_lines(text_lines)
                    if not joined.strip():
                        continue
                    level = 0
                    if max_size >= body_size * 1.6 and len(joined) < 80:
                        level = 1
                    elif max_size >= body_size * 1.25 and len(joined) < 80 and all(_is_bold(sp) for sp in spans_all):
                        level = 2
                    if level:
                        p = d.add_heading(joined, level)
                    else:
                        p = d.add_paragraph()
                        # rebuild runs with formatting, merging spans with identical style
                        cur_style = None
                        buf = ""
                        run_specs: list[tuple[tuple, str]] = []
                        for li, ln in enumerate(self_lines):
                            for sp in ln.get("spans", []):
                                t = sp.get("text", "")
                                if not t:
                                    continue
                                style = (_is_bold(sp), _is_italic(sp), round(float(sp.get("size", body_size)), 1), int(sp.get("color", 0)))
                                if style != cur_style and buf:
                                    run_specs.append((cur_style, buf))
                                    buf = ""
                                cur_style = style
                                buf += t
                            if li < len(self_lines) - 1 and buf and not (_cjk(buf[-1])):
                                buf += " "
                        if buf:
                            run_specs.append((cur_style, buf))
                        for style, text in run_specs:
                            run = p.add_run(text)
                            bold, italic, size, color = style
                            run.bold = bold
                            run.italic = italic
                            if abs(size - body_size) > 0.6:
                                run.font.size = Pt(max(6, min(72, size)))
                            if color not in (0, 0x000000):
                                rr, gg, bb = _rgb(color)
                                if (rr, gg, bb) != (0, 0, 0):
                                    run.font.color.rgb = RGBColor(rr, gg, bb)
                    r = pymupdf.Rect(payload["bbox"])
                    center = (r.x0 + r.x1) / 2
                    if abs(center - page.rect.width / 2) < 12 and r.width < page.rect.width * 0.6 and len(joined) < 60:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif kind == "image":
                    data = payload.get("image")
                    if not data:
                        continue
                    r = pymupdf.Rect(payload["bbox"])
                    w_pt = max(20.0, min(r.width, usable_w))
                    try:
                        d.add_picture(io.BytesIO(data), width=Emu(int(w_pt * _EMU_PER_PT)))
                        n_images += 1
                    except Exception:
                        try:
                            pix = pymupdf.Pixmap(data)
                            if pix.n - pix.alpha >= 4:
                                pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
                            d.add_picture(io.BytesIO(pix.tobytes("png")), width=Emu(int(w_pt * _EMU_PER_PT)))
                            n_images += 1
                        except Exception:
                            warnings.append("部分图片无法嵌入")
                else:
                    data = payload
                    rows = len(data)
                    cols = max(len(rw) for rw in data)
                    if rows == 0 or cols == 0:
                        continue
                    tbl = d.add_table(rows=rows, cols=cols)
                    try:
                        tbl.style = "Table Grid"
                    except Exception:
                        pass
                    for ri, rw in enumerate(data):
                        for ci in range(cols):
                            val = rw[ci] if ci < len(rw) else None
                            tbl.cell(ri, ci).text = "" if val is None else str(val).strip()
                    n_tables += 1
                    d.add_paragraph()
            if pno < n - 1:
                d.add_page_break()
        dst.parent.mkdir(parents=True, exist_ok=True)
        d.save(str(dst))
    finally:
        doc.close()
    warnings.insert(0, f"内置引擎：文字流式重排，{n_tables} 个表格、{n_images} 张图片已保留")
    return warnings
